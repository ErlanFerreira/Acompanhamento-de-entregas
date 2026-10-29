"""Controle de entregas da operação GDC (antes era uma planilha manual): cada
lançamento é uma guia com uma ou mais notas fiscais, destinatário, cidade,
data, manifesto (baixado ou não), devolução e retenção de paletes.

Mora num banco à parte (GDC_DATABASE_URL), com metadata e sessão próprias --
o histórico não disputa a cota de rede do banco principal nem depende dele.
Sem a variável, as tabelas ficam no banco principal (e a tela avisa).
"""

import datetime
import io
import os

import openpyxl
from fastapi import APIRouter, Depends, Request
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Boolean, Column, Date, DateTime, ForeignKey, Integer, String, Text, case, func, inspect, or_, text
from sqlalchemy.orm import Session, declarative_base, relationship, selectinload, sessionmaker

from . import db as banco
from .security import require_acesso_completo_api, usuario_logado

_url_gdc = banco._url(os.environ.get("GDC_DATABASE_URL", "").strip())
BANCO_PROPRIO = bool(_url_gdc)
engine_gdc = banco._criar_engine(_url_gdc) if _url_gdc else banco.engine

GdcBase = declarative_base()
SessionGdc = sessionmaker(bind=engine_gdc, autoflush=False, autocommit=False)

BRASILIA = datetime.timezone(datetime.timedelta(hours=-3))
POR_PAGINA = 100

router = APIRouter()


class GdcEntrega(GdcBase):
    __tablename__ = "gdc_entregas"

    id = Column(Integer, primary_key=True)
    guia = Column(String(40), nullable=False, index=True)
    destinatario = Column(String(200), nullable=False, default="")
    cidade = Column(String(120), nullable=False, default="")
    data = Column(Date, nullable=False, index=True)
    manifesto = Column(String(40), nullable=False, default="", index=True)
    manifesto_baixado = Column(Boolean, nullable=False, default=False)
    devolucao = Column(Boolean, nullable=False, default=False)
    nfd = Column(String(60), nullable=False, default="")  # NF de devolução
    obs_devolucao = Column(Text, nullable=False, default="")
    retencao_palets = Column(Boolean, nullable=False, default=False)
    qtd_palets = Column(Integer, nullable=False, default=0)
    observacao = Column(Text, nullable=False, default="")
    criado_por = Column(String(120), nullable=False, default="")
    criado_em = Column(DateTime, nullable=False)
    atualizado_por = Column(String(120))
    atualizado_em = Column(DateTime)

    notas = relationship("GdcNota", cascade="all, delete-orphan", order_by="GdcNota.id", lazy="selectin")


class GdcNota(GdcBase):
    __tablename__ = "gdc_notas"

    id = Column(Integer, primary_key=True)
    entrega_id = Column(Integer, ForeignKey("gdc_entregas.id", ondelete="CASCADE"), nullable=False, index=True)
    numero = Column(String(20), nullable=False, index=True)


# Colunas criadas depois que a tabela já existia -- `create_all` não altera
# tabelas existentes.
_COLUNAS_ADICIONADAS = {
    "gdc_entregas": {"nfd": "VARCHAR(60) NOT NULL DEFAULT ''", "obs_devolucao": "TEXT NOT NULL DEFAULT ''"},
}

_tabelas_ok = False


def _preparar_tabelas() -> None:
    GdcBase.metadata.create_all(bind=engine_gdc)
    se_nao_existe = "IF NOT EXISTS " if engine_gdc.dialect.name == "postgresql" else ""
    for tabela, colunas in _COLUNAS_ADICIONADAS.items():
        existentes = {c["name"] for c in inspect(engine_gdc).get_columns(tabela)}
        for nome, tipo in colunas.items():
            if nome not in existentes:
                with engine_gdc.begin() as conn:
                    conn.execute(text(f"ALTER TABLE {tabela} ADD COLUMN {se_nao_existe}{nome} {tipo}"))


def get_gdc_db():
    """Sessão do banco GDC. As tabelas são criadas na primeira vez que alguém
    usa a tela (a função da Vercel pode subir sem rodar o seed)."""
    global _tabelas_ok
    if not _tabelas_ok:
        _preparar_tabelas()
        _tabelas_ok = True
    db = SessionGdc()
    try:
        yield db
    finally:
        db.close()


def _agora() -> datetime.datetime:
    return datetime.datetime.now(BRASILIA).replace(tzinfo=None)


def _quem(request: Request) -> str:
    u = usuario_logado(request)
    return (u.get("nome") or u.get("login") or "Analista")[:120] if u else "Supervisor"


def _to_dict(e: GdcEntrega) -> dict:
    return {
        "id": e.id, "guia": e.guia, "notas": [n.numero for n in e.notas],
        "destinatario": e.destinatario, "cidade": e.cidade, "data": e.data.isoformat(),
        "manifesto": e.manifesto, "manifesto_baixado": e.manifesto_baixado,
        "devolucao": e.devolucao, "nfd": e.nfd, "obs_devolucao": e.obs_devolucao, "retencao_palets": e.retencao_palets, "qtd_palets": e.qtd_palets,
        "observacao": e.observacao,
        "criado_por": e.criado_por, "criado_em": e.criado_em.isoformat() if e.criado_em else None,
        "atualizado_por": e.atualizado_por, "atualizado_em": e.atualizado_em.isoformat() if e.atualizado_em else None,
    }


def _sim_nao(valor) -> bool | None:
    return {"sim": True, "nao": False}.get(str(valor or "").lower())


def _filtrar(q, busca: str, de: str, ate: str, baixado: str, devolucao: str, retencao: str):
    busca = busca.strip()
    if busca:
        termo = f"%{busca}%"
        q = q.filter(or_(
            GdcEntrega.guia.ilike(termo), GdcEntrega.manifesto.ilike(termo),
            GdcEntrega.destinatario.ilike(termo), GdcEntrega.cidade.ilike(termo),
            GdcEntrega.observacao.ilike(termo), GdcEntrega.nfd.ilike(termo),
            GdcEntrega.obs_devolucao.ilike(termo), GdcEntrega.notas.any(GdcNota.numero.ilike(termo)),
        ))
    for valor, op in ((de, "__ge__"), (ate, "__le__")):
        try:
            q = q.filter(getattr(GdcEntrega.data, op)(datetime.date.fromisoformat(valor)))
        except ValueError:
            pass
    for valor, coluna in ((baixado, GdcEntrega.manifesto_baixado), (devolucao, GdcEntrega.devolucao),
                          (retencao, GdcEntrega.retencao_palets)):
        if _sim_nao(valor) is not None:
            q = q.filter(coluna.is_(_sim_nao(valor)))
    return q


def _validar(payload: dict) -> tuple[dict | None, str | None]:
    def texto(chave, limite):
        return " ".join(str(payload.get(chave) or "").split())[:limite]

    dados = {
        "guia": texto("guia", 40), "destinatario": texto("destinatario", 200), "cidade": texto("cidade", 120),
        "manifesto": texto("manifesto", 40), "observacao": str(payload.get("observacao") or "").strip()[:4000],
        "manifesto_baixado": bool(payload.get("manifesto_baixado")), "devolucao": bool(payload.get("devolucao")),
        "retencao_palets": bool(payload.get("retencao_palets")),
    }
    # NFD e a observação da devolução só valem quando houve devolução.
    dados["nfd"] = texto("nfd", 60) if dados["devolucao"] else ""
    dados["obs_devolucao"] = str(payload.get("obs_devolucao") or "").strip()[:4000] if dados["devolucao"] else ""
    if not dados["guia"]:
        return None, "Informe o número da guia."
    if not dados["destinatario"]:
        return None, "Informe o destinatário."
    try:
        dados["data"] = datetime.date.fromisoformat(str(payload.get("data") or ""))
    except ValueError:
        return None, "Informe a data."
    try:
        dados["qtd_palets"] = int(payload.get("qtd_palets") or 0) if dados["retencao_palets"] else 0
    except (TypeError, ValueError):
        return None, "Quantidade de paletes inválida."
    if dados["retencao_palets"] and not 0 < dados["qtd_palets"] <= 9999:
        return None, "Houve retenção de paletes: informe a quantidade."
    notas = []
    for n in payload.get("notas") or []:
        n = "".join(str(n).split())[:20]
        if n and n not in notas:
            notas.append(n)
    dados["notas"] = notas
    return dados, None


def _guia_repetida(db: Session, guia: str, ignorar_id: int | None = None) -> GdcEntrega | None:
    q = db.query(GdcEntrega).filter(func.upper(GdcEntrega.guia) == guia.upper())
    if ignorar_id:
        q = q.filter(GdcEntrega.id != ignorar_id)
    return q.first()


def _gravar(db: Session, e: GdcEntrega, dados: dict) -> None:
    for chave, valor in dados.items():
        if chave != "notas":
            setattr(e, chave, valor)
    e.notas = [GdcNota(numero=n) for n in dados["notas"]]


def _conta(condicao):
    return func.coalesce(func.sum(case((condicao, 1), else_=0)), 0)


@router.get("/api/gdc", dependencies=[Depends(require_acesso_completo_api)])
def api_gdc_lista(busca: str = "", de: str = "", ate: str = "", baixado: str = "", devolucao: str = "",
                  retencao: str = "", pagina: int = 1, db: Session = Depends(get_gdc_db)):
    q = _filtrar(db.query(GdcEntrega), busca, de, ate, baixado, devolucao, retencao)
    total, nao_baixados, devolucoes, retencoes, palets = _filtrar(db.query(
        func.count(GdcEntrega.id),
        _conta(GdcEntrega.manifesto_baixado.is_(False)),
        _conta(GdcEntrega.devolucao.is_(True)),
        _conta(GdcEntrega.retencao_palets.is_(True)),
        func.coalesce(func.sum(GdcEntrega.qtd_palets), 0),
    ), busca, de, ate, baixado, devolucao, retencao).one()
    pagina = max(1, pagina)
    itens = (q.options(selectinload(GdcEntrega.notas))
             .order_by(GdcEntrega.data.desc(), GdcEntrega.id.desc())
             .offset((pagina - 1) * POR_PAGINA).limit(POR_PAGINA).all())
    return {
        "itens": [_to_dict(e) for e in itens], "pagina": pagina, "por_pagina": POR_PAGINA,
        "resumo": {"total": int(total), "nao_baixados": int(nao_baixados), "devolucoes": int(devolucoes),
                   "retencoes": int(retencoes), "palets": int(palets)},
    }


@router.post("/api/gdc", dependencies=[Depends(require_acesso_completo_api)])
def api_gdc_criar(request: Request, payload: dict, db: Session = Depends(get_gdc_db)):
    dados, erro = _validar(payload)
    if erro:
        return JSONResponse({"erro": erro}, status_code=400)
    repetida = _guia_repetida(db, dados["guia"])
    if repetida and not payload.get("confirmar_repetida"):
        return JSONResponse({"erro": f"A guia {repetida.guia} já foi lançada em {repetida.data:%d/%m/%Y}.",
                             "repetida": True}, status_code=409)
    e = GdcEntrega(criado_por=_quem(request), criado_em=_agora())
    _gravar(db, e, dados)
    db.add(e)
    db.commit()
    return _to_dict(e)


@router.put("/api/gdc/{entrega_id}", dependencies=[Depends(require_acesso_completo_api)])
def api_gdc_editar(request: Request, entrega_id: int, payload: dict, db: Session = Depends(get_gdc_db)):
    e = db.get(GdcEntrega, entrega_id)
    if not e:
        return JSONResponse({"erro": "Lançamento não encontrado (pode ter sido excluído)."}, status_code=404)
    dados, erro = _validar(payload)
    if erro:
        return JSONResponse({"erro": erro}, status_code=400)
    # Só avisa se a guia mudou -- uma repetida já confirmada não pergunta de novo a cada edição.
    mudou = dados["guia"].upper() != e.guia.upper()
    repetida = mudou and _guia_repetida(db, dados["guia"], ignorar_id=e.id)
    if repetida and not payload.get("confirmar_repetida"):
        return JSONResponse({"erro": f"A guia {repetida.guia} já foi lançada em {repetida.data:%d/%m/%Y}.",
                             "repetida": True}, status_code=409)
    _gravar(db, e, dados)
    e.atualizado_por, e.atualizado_em = _quem(request), _agora()
    db.commit()
    return _to_dict(e)


@router.delete("/api/gdc/{entrega_id}", dependencies=[Depends(require_acesso_completo_api)])
def api_gdc_excluir(entrega_id: int, db: Session = Depends(get_gdc_db)):
    e = db.get(GdcEntrega, entrega_id)
    if e:
        db.delete(e)
        db.commit()
    return {"ok": True}


@router.get("/api/gdc/exportar", dependencies=[Depends(require_acesso_completo_api)])
def api_gdc_exportar(busca: str = "", de: str = "", ate: str = "", baixado: str = "", devolucao: str = "",
                     retencao: str = "", db: Session = Depends(get_gdc_db)):
    itens = (_filtrar(db.query(GdcEntrega), busca, de, ate, baixado, devolucao, retencao)
             .options(selectinload(GdcEntrega.notas))
             .order_by(GdcEntrega.data.desc(), GdcEntrega.id.desc()).all())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controle GDC"
    cabecalho = ["Data", "Guia", "Notas fiscais", "Destinatário", "Cidade", "Manifesto", "Manifesto baixado",
                 "Devolução", "NFD", "Obs. da devolução", "Retenção de paletes", "Qtd. paletes", "Observação", "Lançado por", "Lançado em"]
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")
    sn = {True: "Sim", False: "Não"}
    for e in itens:
        ws.append([e.data, e.guia, ", ".join(n.numero for n in e.notas), e.destinatario, e.cidade, e.manifesto,
                   sn[e.manifesto_baixado], sn[e.devolucao], e.nfd, e.obs_devolucao, sn[e.retencao_palets], e.qtd_palets or None,
                   e.observacao, e.criado_por, e.criado_em])
        ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
        ws.cell(ws.max_row, 15).number_format = "DD/MM/YYYY HH:MM"
    for letra, largura in zip("ABCDEFGHIJKLMNO", (12, 14, 26, 36, 22, 14, 10, 10, 14, 40, 10, 10, 50, 20, 17)):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"controle_gdc_{datetime.datetime.now(BRASILIA):%Y-%m-%d}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nome}"'})
