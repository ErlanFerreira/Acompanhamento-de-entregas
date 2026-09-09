import datetime
import io
import json
import logging
import re
from contextlib import asynccontextmanager

import openpyxl
import requests
from fastapi import Depends, FastAPI, Form, Request, UploadFile
from fastapi.exception_handlers import http_exception_handler
from fastapi.exceptions import HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from . import config
from .db import get_db
from .models import Carga, ConsultaItem, ConsultaJob, Meta
from .security import (
    COOKIE_MAX_AGE,
    COOKIE_NAME,
    create_session_cookie,
    is_authenticated,
    require_auth_api,
    require_auth_page,
)
from .seed import init_db_and_seed

logging.basicConfig(level=logging.INFO)


@asynccontextmanager
async def lifespan(app: FastAPI):
    # A sincronização periódica roda via GitHub Actions (ver .github/workflows/sync.yml),
    # não dentro do processo web -- necessário porque a Vercel roda o app como função
    # serverless (sem processo de fundo de longa duração).
    init_db_and_seed()
    yield


app = FastAPI(lifespan=lifespan)
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")


@app.exception_handler(HTTPException)
async def redirect_aware_exception_handler(request: Request, exc: HTTPException):
    location = (exc.headers or {}).get("Location")
    if exc.status_code == 303 and location:
        return RedirectResponse(url=location, status_code=303)
    return await http_exception_handler(request, exc)


# ---------------------------------------------------------------- auth ----

@app.get("/", response_class=HTMLResponse)
def root(request: Request):
    return RedirectResponse(url="/painel" if is_authenticated(request) else "/login")


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if is_authenticated(request):
        return RedirectResponse(url="/painel")
    return templates.TemplateResponse("login.html", {"request": request, "error": None})


@app.post("/login", response_class=HTMLResponse)
def login_submit(request: Request, senha: str = Form(...)):
    if not config.PAINEL_SENHA or senha != config.PAINEL_SENHA:
        return templates.TemplateResponse(
            "login.html",
            {"request": request, "error": "Senha incorreta."},
            status_code=401,
        )
    token = create_session_cookie()
    resp = RedirectResponse(url="/painel", status_code=303)
    resp.set_cookie(COOKIE_NAME, token, max_age=COOKIE_MAX_AGE, httponly=True, samesite="lax")
    return resp


@app.get("/logout")
def logout():
    resp = RedirectResponse(url="/login", status_code=303)
    resp.delete_cookie(COOKIE_NAME)
    return resp


# ------------------------------------------------------------- painel -----

@app.get("/painel", response_class=HTMLResponse, dependencies=[Depends(require_auth_page)])
def painel_page(request: Request):
    return templates.TemplateResponse("painel.html", {"request": request, "active": "painel"})


def _row_to_dict(r: Carga) -> dict:
    return {
        "emissao": r.emissao,
        "cte": r.cte,
        "remetente": r.remetente,
        "consignatario": r.consignatario,
        "cidade_cons": r.cidade_cons,
        "destinatario": r.destinatario,
        "cidade_dest": r.cidade_dest,
        "uf_dest": r.uf_dest,
        "manifesto": r.manifesto,
        "romaneio": r.romaneio,
        "ocorrencia": r.ocorrencia,
        "entregue": r.entregue,
        "data_comprovante": r.data_comprovante,
        "data_baixa": r.data_baixa,
        "status": r.status,
        "estado": r.estado,
        "previsao": r.previsao,
        "dias_atraso": r.dias_atraso,
    }


@app.get("/api/cargas", dependencies=[Depends(require_auth_api)])
def api_cargas(
    data_inicial: str | None = None,
    data_final: str | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(Carga)
    if data_inicial:
        q = q.filter(Carga.emissao >= data_inicial)
    if data_final:
        q = q.filter(Carga.emissao <= data_final)
    rows = q.all()
    records = [_row_to_dict(r) for r in rows]

    last_sync_at = db.query(Meta).filter(Meta.chave == "last_sync_at").first()
    gerado_em = last_sync_at.valor[:10] if last_sync_at and last_sync_at.valor else datetime.date.today().isoformat()

    meta = {
        "gerado_em": gerado_em,
        "referencia": datetime.date.today().isoformat(),
        "total": len(records),
        "fonte": "Webtrans (GW Sistemas) - relatório \"Pendências\", atualização automática",
    }
    return {"meta": meta, "records": records}


def _dispatch_workflow(workflow_file: str, inputs: dict | None = None):
    """Dispara um workflow do GitHub Actions (workflow_dispatch).

    O app web (Vercel) não roda a automação de navegador diretamente -- ela
    precisa de um Chromium instalado, inviável numa função serverless. Isso
    só enfileira o job no GitHub; quem processa é o runner do Actions.
    """
    if not config.GITHUB_REPO or not config.GITHUB_DISPATCH_TOKEN:
        return {"ok": False, "erro": "GITHUB_REPO/GITHUB_DISPATCH_TOKEN não configurados nesta implantação."}
    url = f"https://api.github.com/repos/{config.GITHUB_REPO}/actions/workflows/{workflow_file}/dispatches"
    headers = {
        "Authorization": f"Bearer {config.GITHUB_DISPATCH_TOKEN}",
        "Accept": "application/vnd.github+json",
    }
    body = {"ref": "main"}
    if inputs:
        body["inputs"] = inputs
    try:
        resp = requests.post(url, headers=headers, json=body, timeout=15)
        if resp.status_code >= 300:
            return {"ok": False, "erro": f"GitHub respondeu {resp.status_code}: {resp.text[:300]}"}
        return {"ok": True}
    except requests.RequestException as exc:
        return {"ok": False, "erro": str(exc)}


@app.post("/api/sync/run", dependencies=[Depends(require_auth_api)])
def api_sync_run():
    """Dispara a sincronização periódica fora do horário agendado.

    Assíncrono -- o resultado aparece no painel após o workflow terminar
    (~1-2 min), não instantaneamente.
    """
    resultado = _dispatch_workflow("sync.yml")
    if not resultado["ok"]:
        return JSONResponse(resultado, status_code=502 if "GitHub respondeu" in resultado.get("erro", "") else 500)
    return {"ok": True, "queued": True}


# --------------------------------------------------- consulta de notas ----

@app.get("/consultas", response_class=HTMLResponse, dependencies=[Depends(require_auth_page)])
def consultas_page(request: Request):
    return templates.TemplateResponse("consultas.html", {"request": request, "active": "consultas"})


def _job_to_dict(job: ConsultaJob) -> dict:
    return {
        "id": job.id,
        "criado_em": job.criado_em.isoformat() if job.criado_em else None,
        "nome_arquivo": job.nome_arquivo,
        "status": job.status,
        "total_itens": job.total_itens,
        "processados": job.processados,
        "erro_mensagem": job.erro_mensagem,
    }


@app.get("/api/consultas", dependencies=[Depends(require_auth_api)])
def api_consultas_list(db: Session = Depends(get_db)):
    jobs = db.query(ConsultaJob).order_by(ConsultaJob.criado_em.desc()).limit(30).all()
    return {"jobs": [_job_to_dict(j) for j in jobs]}


@app.get("/api/consultas/{job_id}", dependencies=[Depends(require_auth_api)])
def api_consultas_status(job_id: int, db: Session = Depends(get_db)):
    job = db.query(ConsultaJob).filter(ConsultaJob.id == job_id).first()
    if not job:
        return JSONResponse({"erro": "Job não encontrado."}, status_code=404)
    return _job_to_dict(job)


def _primeira_nf(valor) -> str:
    """Às vezes a célula tem mais de uma NF (mesmo conhecimento, notas
    diferentes), ex.: "59185 / 59186" -- usa só a primeira para a consulta
    (o valor original completo continua preservado em `linha_original`)."""
    texto = str(valor).strip()
    return re.split(r"[/,;-]", texto)[0].strip()


@app.post("/api/consultas", dependencies=[Depends(require_auth_api)])
async def api_consultas_create(
    arquivo: UploadFile,
    coluna_nf: str = Form(...),
    db: Session = Depends(get_db),
):
    conteudo = await arquivo.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    except Exception:
        return JSONResponse({"erro": "Não consegui ler o arquivo. Confirme que é um .xlsx válido."}, status_code=400)

    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return JSONResponse({"erro": "Planilha vazia."}, status_code=400)

    cabecalho = [str(c).strip() if c is not None else "" for c in linhas[0]]
    if coluna_nf not in cabecalho:
        return JSONResponse({"erro": f'Coluna "{coluna_nf}" não encontrada na planilha.'}, status_code=400)
    idx_nf = cabecalho.index(coluna_nf)

    job = ConsultaJob(nome_arquivo=arquivo.filename, coluna_nf=coluna_nf, status="pendente")
    db.add(job)
    db.flush()

    total = 0
    for i, row in enumerate(linhas[1:], start=1):
        if row is None or all(c is None for c in row):
            continue
        nf_valor = row[idx_nf] if idx_nf < len(row) else None
        if nf_valor is None or str(nf_valor).strip() == "":
            continue
        linha_dict = {cabecalho[c]: row[c] for c in range(len(cabecalho)) if cabecalho[c] and c < len(row)}
        item = ConsultaItem(
            job_id=job.id,
            linha_idx=i,
            linha_original=json.dumps(linha_dict, ensure_ascii=False, default=str),
            nf_numero=_primeira_nf(nf_valor),
        )
        db.add(item)
        total += 1

    if total == 0:
        db.rollback()
        return JSONResponse({"erro": f'Nenhuma linha com valor preenchido na coluna "{coluna_nf}".'}, status_code=400)

    job.total_itens = total
    db.commit()

    resultado = _dispatch_workflow("consulta.yml", inputs={"job_id": str(job.id)})
    if not resultado["ok"]:
        job.status = "erro"
        job.erro_mensagem = resultado["erro"][:490]
        db.commit()
        return JSONResponse({"erro": resultado["erro"]}, status_code=502)

    return {"ok": True, "job_id": job.id, "total_itens": total}


@app.get("/api/consultas/{job_id}/arquivo", dependencies=[Depends(require_auth_api)])
def api_consultas_arquivo(job_id: int, db: Session = Depends(get_db)):
    job = db.query(ConsultaJob).filter(ConsultaJob.id == job_id).first()
    if not job:
        return JSONResponse({"erro": "Job não encontrado."}, status_code=404)

    itens = (
        db.query(ConsultaItem)
        .filter(ConsultaItem.job_id == job_id)
        .order_by(ConsultaItem.linha_idx.asc(), ConsultaItem.id.asc())
        .all()
    )

    cabecalho_original = list(json.loads(itens[0].linha_original).keys()) if itens else []
    colunas_novas = ["CT-e", "Status Entrega", "Previsão Entrega", "Data Entrega", "Dias Atraso", "Observação", "Encontrado?"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultado"
    ws.append(cabecalho_original + colunas_novas)

    for item in itens:
        linha = json.loads(item.linha_original)
        valores_originais = [linha.get(c) for c in cabecalho_original]
        valores_novos = [
            item.cte,
            item.status_entrega,
            item.previsao_entrega,
            item.data_entrega,
            item.dias_atraso,
            item.observacao,
            "Sim" if item.encontrado else "Não",
        ]
        ws.append(valores_originais + valores_novos)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_saida = f"consulta_{job_id}_resultado.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome_saida}"'},
    )
