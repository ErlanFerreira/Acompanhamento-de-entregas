"""Relatórios gerados a partir da tela de Bipagem, preenchendo os modelos em
`app/modelos/` (cópias limpas das planilhas que o SAC já usava -- mantêm
logo, cores, fórmulas e layout de impressão):

- Controle de comprovantes (`controle_comprovantes.xlsx`): todos os CT-e do
  mês de um tomador; os bipados ganham a data do (primeiro) bip em
  "Data de Envio ao Financeiro/Arquivo" e viram "Recebido".
- Protocolo de envio de faturas (`protocolo_envio.xlsx`): lista "DACTE / NF"
  dos CT-e bipados, uma aba por filial + tomador (o recebedor).
"""

import copy
import datetime
import io
import re
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.formatting.formatting import ConditionalFormattingList
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from . import chave_cte
from .escopo import carga_visivel, filtro_cargas
from .filiais import dados_filial, fmt_cnpj
from .models import Bipagem, Carga, ProtocoloFatura

MODELOS = Path(__file__).parent / "modelos"
# Brasília sem horário de verão (abolido em 2019) -- evita depender do
# pacote tzdata, que não vem no Windows.
FUSO_BR = datetime.timezone(datetime.timedelta(hours=-3))

# Um bip serve pros dois documentos. Registros antigos podem ter
# "comprovante"/"fatura" (quando as listas eram separadas) -- as consultas
# ignoram a finalidade.
FINALIDADE_GERAL = "geral"


def data_br(dt_utc: datetime.datetime) -> datetime.date:
    return dt_utc.replace(tzinfo=datetime.timezone.utc).astimezone(FUSO_BR).date()


def _serie_confere(carga: Carga, dados: dict) -> bool:
    # Série só vem se o relatório do GW tiver essa coluna. Minuta no GW é
    # série "M"; na chave impressa a série é numérica, então aceita os dois.
    if not carga.serie:
        return True
    serie = carga.serie.strip().upper()
    if dados["modelo"] == chave_cte.MODELO_MINUTA:
        return serie in ("M", dados["serie"])
    return serie == dados["serie"]


def buscar_carga(db: Session, dados: dict) -> Carga | None:
    """Acha o CT-e/minuta da chave decodificada no banco -- mesma filial e
    número; havendo mais de um (CT-e e minuta com o mesmo número), a série
    decide quando o relatório a traz."""
    # No banco o número vem como no relatório do GW (ex: "088375", 6 dígitos
    # com zeros à esquerda); na chave são 9 dígitos -- testa as variações.
    numero = dados["numero"]
    candidatos = {str(numero)} | {str(numero).zfill(w) for w in range(6, 10)}
    cargas = (
        db.query(Carga)
        .filter(Carga.cnpj_filial == dados["cnpj_emitente"], Carga.cte.in_(candidatos))
        .all()
    )
    # Os que têm série (e ela confere) antes dos sem série, gravados antes
    # do relatório trazer essa coluna.
    cargas.sort(key=lambda c: c.serie is None)
    return next((c for c in cargas if _serie_confere(c, dados)), None)


def serie_normalizada(serie: str | None) -> str | None:
    """Série comparável entre bip e CT-e: a minuta é "M" no GW e "99" na
    chave impressa."""
    s = str(serie or "").strip().upper()
    if not s:
        return None
    return "M" if s in ("M", "99") else s.lstrip("0") or "0"


def _buscar_por_redespacho(db: Session, dados: dict) -> Carga | None:
    """Chave do DACTE do PARCEIRO: o número dela é o "N. CT-e Redespacho" de
    um CT-e nosso, e o emitente (o parceiro) é o tomador desse CT-e -- a raiz
    do CNPJ precisa bater, pra não confundir parceiros diferentes que usam o
    mesmo número."""
    numero = dados["numero"]
    candidatos = {str(numero)} | {str(numero).zfill(w) for w in range(6, 10)}
    raiz_parceiro = dados["cnpj_emitente"][:8]
    cargas = [
        c for c in db.query(Carga).filter(Carga.cte_redespacho.in_(candidatos))
        if (c.cnpj_consignatario or "")[:8] == raiz_parceiro
    ]
    # Mais de um (número reaproveitado pelo parceiro em outro período): o
    # mais recente.
    cargas.sort(key=lambda c: c.emissao or "", reverse=True)
    return cargas[0] if cargas else None


def _numeros_nf(texto: str | None) -> set[int]:
    """NFs da coluna "Notas" (ex: "589002,589003")."""
    return {int(t) for t in re.findall(r"\d+", texto or "")}


def buscar_por_nf(db: Session, numero: int, cnpj_emitente: str | None = None) -> list[Carga]:
    """CT-e/minutas nossos que transportaram a NF.

    Com o CNPJ do emitente da NF (vem na chave da DANFE), só os CT-e em que
    a raiz dele é a do remetente (ou do tomador) -- o mesmo número de NF
    existe em empresas diferentes. CT-e sincronizados antes de existir o
    CNPJ do remetente na base só entram se nenhum CT-e conferir."""
    cargas = [
        c for c in db.query(Carga).filter(Carga.notas_fiscais.like(f"%{numero}%"))
        if numero in _numeros_nf(c.notas_fiscais)
    ]
    if cnpj_emitente:
        raiz = cnpj_emitente[:8]
        conferem = [
            c for c in cargas
            if raiz in ((c.cnpj_remetente or "")[:8], (c.cnpj_consignatario or "")[:8])
        ]
        cargas = conferem or [c for c in cargas if not c.cnpj_remetente]
    cargas.sort(key=lambda c: c.emissao or "", reverse=True)
    return cargas


def _dados_do_nosso_documento(dados: dict, carga: Carga) -> dict:
    """Dados da leitura trocados pela identidade do NOSSO CT-e/minuta
    (filial, número, série, tipo) -- quando o que foi bipado é a chave de
    outro documento (DACTE do parceiro, DANFE) ou o número da NF."""
    minuta = serie_normalizada(carga.serie) == "M"
    modelo = chave_cte.MODELO_MINUTA if minuta else chave_cte.MODELO_CTE
    return {
        **dados,
        "cnpj_emitente": carga.cnpj_filial,
        "cnpj_emitente_fmt": fmt_cnpj(carga.cnpj_filial),
        "numero": int(carga.cte) if (carga.cte or "").isdigit() else dados.get("numero"),
        "serie": "M" if minuta else (carga.serie or ""),
        "modelo": modelo,
        "tipo": chave_cte.TIPOS[modelo],
    }


def _resolver_nf(db: Session, numero: int, cnpj_emitente: str | None, dados: dict) -> tuple[dict, Carga | None]:
    cargas = buscar_por_nf(db, numero, cnpj_emitente)
    base = {**dados, "via_nf": True, "nf_bipada": numero}
    if not cargas:
        return {**base, "nf_nao_encontrada": True}, None
    if len({c.id_cte for c in cargas}) > 1:
        # NF em mais de um CT-e (ou, sem a DANFE, de empresas diferentes) --
        # não dá pra escolher sozinho.
        base["nf_ambigua"] = [
            {"cte": c.cte, "serie": c.serie, "remetente": c.remetente, "emissao": c.emissao}
            for c in cargas[:8]
        ]
        return base, None
    carga = cargas[0]
    nossos = _dados_do_nosso_documento(base, carga)
    if not nossos.get("chave") or len(nossos["chave"]) != 44:
        # Número de NF digitado: não há chave -- o bip é identificado pela
        # NF + o nosso documento.
        nossos["chave"] = f"NF:{numero}:{carga.cnpj_filial}:{nossos['numero']}:{nossos['serie']}"
    return nossos, carga


def resolver_chave(db: Session, dados: dict) -> tuple[dict, Carga | None]:
    """Acha o NOSSO CT-e/minuta de uma chave bipada -- a nossa chave ou a do
    DACTE do parceiro (redespacho). Devolve os dados da chave já com a
    identidade do nosso documento (filial, número, série, tipo) e a carga.

    Bipada pela chave do parceiro, vem também `via_parceiro=True` e
    `cte_parceiro_bipado` (o número dele); `chave` continua sendo a chave
    lida, que identifica o bip."""
    if dados["modelo"] == chave_cte.MODELO_NFE:
        return _resolver_nf(db, dados["numero"], dados["cnpj_emitente"], dados)
    carga = buscar_carga(db, dados)
    if carga or dados["modelo"] == chave_cte.MODELO_MINUTA:
        return dados, carga
    carga = _buscar_por_redespacho(db, dados)
    if not carga:
        return dados, None
    nossos = _dados_do_nosso_documento(dados, carga)
    nossos.update({
        "via_parceiro": True,
        "cte_parceiro_bipado": dados["numero"],
        "cnpj_parceiro_fmt": dados["cnpj_emitente_fmt"],
    })
    return nossos, carga


def resolver_leitura(db: Session, leitura: str) -> tuple[dict, Carga | None]:
    """`resolver_chave` a partir do texto lido: chave de 44 dígitos (ou QR
    Code), ou só o número de uma NF digitado (até 9 dígitos). Levanta
    `ChaveInvalida` se não for nenhum dos dois."""
    texto = (leitura or "").strip()
    if re.fullmatch(r"\d{1,9}", texto) and int(texto) > 0:
        return _resolver_nf(db, int(texto), None, {"chave": None, "numero": int(texto)})
    return resolver_chave(db, chave_cte.decodificar(texto))


def carga_do_bip(db: Session, bip: Bipagem) -> Carga | None:
    """CT-e/minuta de um bip gravado: pela identidade do nosso documento
    (filial + número + série, gravada em todo bip); se o documento ainda não
    estava na base quando foi bipado, refaz a busca pela chave lida."""
    serie = serie_normalizada(bip.serie)
    minuta = serie == "M"
    carga = buscar_carga(db, {
        "cnpj_emitente": bip.cnpj_filial,
        "numero": bip.numero,
        "serie": "99" if minuta else (serie or ""),
        "modelo": chave_cte.MODELO_MINUTA if minuta else chave_cte.MODELO_CTE,
    })
    if carga:
        return carga
    try:
        return resolver_leitura(db, bip.chave)[1]
    except chave_cte.ChaveInvalida:
        return None


def _data(iso: str | None):
    return datetime.date.fromisoformat(iso) if iso else None


def _int_se_numero(v):
    texto = str(v or "").strip()
    return int(texto) if texto.isdigit() else (texto or None)


def _salvar(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return preservar_espacos(buffer.getvalue())


def preservar_espacos(xlsx: bytes) -> bytes:
    """O openpyxl grava trechos de rich text só com espaços (a linha da
    assinatura no rodapé do protocolo) como `<t>   </t>`, sem
    `xml:space="preserve"` -- o Excel lê isso como trecho vazio e acusa
    "Encontramos um problema em um conteúdo" ao abrir. Marca todos os `<t>`
    pra preservar os espaços."""
    entrada = zipfile.ZipFile(io.BytesIO(xlsx))
    saida_buffer = io.BytesIO()
    with zipfile.ZipFile(saida_buffer, "w", zipfile.ZIP_DEFLATED) as saida:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") or item.filename == "xl/sharedStrings.xml":
                dados = dados.replace(b"<t>", b'<t xml:space="preserve">')
            saida.writestr(item, dados)
    return saida_buffer.getvalue()


# ------------------------------------------------ controle de comprovantes ----

_LINHA_INICIAL = 9
_NUM_COLUNAS = 18  # A..R


def gerar_controle_comprovantes(
    db: Session, tomador: str, ano: int, mes: int, responsavel: str, empresas: set[str] | None = None
) -> bytes:
    """`tomador` é o `id` de `tomadores_disponiveis`: raiz do CNPJ (8
    dígitos) ou "nome:<nome>" pra CT-e ainda sem CNPJ do tomador. Com
    `empresas` (analista), só entram os CT-e dessas empresas."""
    inicio = datetime.date(ano, mes, 1)
    fim = datetime.date(ano + (mes == 12), mes % 12 + 1, 1)
    if tomador.startswith(PREFIXO_TOMADOR_NOME):
        nome = tomador[len(PREFIXO_TOMADOR_NOME):]
        filtro_tomador = Carga.consignatario == nome
    else:
        # CT-e ainda sem CNPJ do tomador entram pelo nome, se for o mesmo
        # nome que esse CNPJ tem nos CT-e que já têm.
        nomes = [
            n for (n,) in db.query(Carga.consignatario)
            .filter(Carga.cnpj_consignatario.like(f"{tomador}%"))
            .distinct()
            if n
        ]
        filtro_tomador = or_(
            Carga.cnpj_consignatario.like(f"{tomador}%"),
            and_(Carga.cnpj_consignatario.is_(None), Carga.consignatario.in_(nomes)),
        )
    q = db.query(Carga).filter(
        filtro_tomador,
        Carga.emissao >= inicio.isoformat(),
        Carga.emissao < fim.isoformat(),
    )
    if empresas is not None:
        q = q.filter(filtro_cargas(empresas))
    cargas = q.order_by(Carga.emissao.asc(), Carga.cte.asc()).all()

    # Primeiro bip de cada CT-e. CT-e e minuta usam a mesma numeração na
    # mesma filial, então casa pela série também; CT-e sem série na base
    # (sincronizado antes dessa coluna) casa só por filial + número.
    bips_por_serie: dict[tuple[str, int, str | None], Bipagem] = {}
    bips_qualquer: dict[tuple[str, int], Bipagem] = {}
    filiais = {c.cnpj_filial for c in cargas}
    if filiais:
        for b in (
            db.query(Bipagem)
            .filter(Bipagem.cnpj_filial.in_(filiais))
            .order_by(Bipagem.bipado_em.asc())
        ):
            bips_por_serie.setdefault((b.cnpj_filial, b.numero, serie_normalizada(b.serie)), b)
            bips_qualquer.setdefault((b.cnpj_filial, b.numero), b)

    def bip_do_cte(carga: Carga, numero: int | None) -> Bipagem | None:
        serie = serie_normalizada(carga.serie)
        if serie is None:
            return bips_qualquer.get((carga.cnpj_filial, numero))
        return bips_por_serie.get((carga.cnpj_filial, numero, serie)) or bips_por_serie.get(
            (carga.cnpj_filial, numero, None)
        )

    wb = openpyxl.load_workbook(MODELOS / "controle_comprovantes.xlsx", rich_text=True)
    ws = wb.active
    ws["H4"] = inicio
    ws["H5"] = responsavel or None

    estilos = [ws.cell(row=_LINHA_INICIAL, column=c) for c in range(1, _NUM_COLUNAS + 1)]
    for i, carga in enumerate(cargas):
        r = _LINHA_INICIAL + i
        numero = int(carga.cte) if (carga.cte or "").isdigit() else None
        bip = bip_do_cte(carga, numero)
        valores = [
            carga.cte,
            carga.serie or (bip.serie if bip else None),
            carga.tipo_conhecimento,
            carga.filial or dados_filial(carga.cnpj_filial)["apelido"],
            _data(carga.emissao),
            _data(carga.previsao),
            fmt_cnpj(carga.cnpj_consignatario),
            carga.consignatario,
            carga.remetente,
            carga.destinatario,
            carga.cidade_dest,
            carga.uf_dest,
            carga.notas_fiscais,
            carga.cte_redespacho,
            data_br(bip.bipado_em) if bip else None,
            f'=IF(ISBLANK($O{r}), "Pendente", "Recebido")',
            None,
            None,
        ]
        for c, (valor, estilo) in enumerate(zip(valores, estilos), start=1):
            celula = ws.cell(row=r, column=c, value=valor)
            if r != _LINHA_INICIAL:
                celula.font = copy.copy(estilo.font)
                celula.border = copy.copy(estilo.border)
                celula.fill = copy.copy(estilo.fill)
                celula.alignment = copy.copy(estilo.alignment)
                celula.number_format = estilo.number_format

    ultima = _LINHA_INICIAL + max(len(cargas), 1) - 1
    ws.auto_filter.ref = f"A8:R{ultima}"

    # A formatação condicional do modelo (verde "Recebido" / vermelho
    # "Pendente") cobre um intervalo fixo -- refaz pro tamanho real.
    regras = [regra for cf in ws.conditional_formatting for regra in cf.rules]
    ws.conditional_formatting = ConditionalFormattingList()
    for regra in regras:
        texto = regra.text or ""
        regra.formula = [f'NOT(ISERROR(SEARCH("{texto}",P{_LINHA_INICIAL})))']
        ws.conditional_formatting.add(f"P{_LINHA_INICIAL}:P{ultima}", regra)

    return _salvar(wb)


PREFIXO_TOMADOR_NOME = "nome:"


def tomadores_disponiveis(db: Session, empresas: set[str] | None = None) -> list[dict]:
    """Tomadores com CT-e na base, pro seletor do controle mensal --
    agrupados pela raiz do CNPJ (matriz + filiais do parceiro). CT-e ainda
    sem CNPJ do tomador (sincronizados antes dessa coluna existir) entram
    pelo nome, com `id` "nome:<nome>". Com `empresas` (analista), só os
    tomadores de CT-e dessas empresas."""
    q = db.query(Carga.cnpj_consignatario, Carga.consignatario, func.count())
    if empresas is not None:
        q = q.filter(filtro_cargas(empresas))
    linhas = q.group_by(Carga.cnpj_consignatario, Carga.consignatario).all()
    grupos: dict[str, dict] = {}
    nomes_com_cnpj: set[str] = set()
    for cnpj, nome, total in linhas:
        if cnpj:
            raiz = cnpj[:8]
            g = grupos.setdefault(raiz, {"id": raiz, "cnpj_raiz": raiz, "nome": nome, "total": 0})
            g["total"] += total
            nomes_com_cnpj.add((nome or "").strip().upper())
    for cnpj, nome, total in linhas:
        nome_norm = (nome or "").strip().upper()
        if cnpj or not nome_norm or nome_norm in nomes_com_cnpj:
            continue
        chave = PREFIXO_TOMADOR_NOME + nome.strip()
        g = grupos.setdefault(chave, {"id": chave, "cnpj_raiz": None, "nome": nome.strip(), "total": 0})
        g["total"] += total
    return sorted(grupos.values(), key=lambda g: (g["nome"] or "").upper())


# ------------------------------------------------ protocolo de faturas ----

# O modelo tem uma página por aba, com 41 linhas (10 a 50; a 51 é só um
# espaçador) e dois pares "DACTE / NF": colunas A (mesclada com B) / C e
# D / E. Preenche o primeiro par de cima a baixo, depois o segundo.
# Cada posição: (linha, coluna do DACTE, coluna da NF).
_LINHAS_PROTOCOLO = range(10, 51)
_SLOTS_PROTOCOLO = [(r, 1, 3) for r in _LINHAS_PROTOCOLO] + [(r, 4, 5) for r in _LINHAS_PROTOCOLO]
# Como o emitente aparece no modelo (que foi feito pra matriz).
_EMITENTE_MODELO = {
    "razao_social": "G. FREIRE BEZERRA DE MORAIS",
    "cnpj": "20.343.618/0001-61",
    "endereco": "R. Jose Jaime Coutinho Dias, 705 a 715 - Bairro Novo / Carpina",
}


def _substituir(celula, antigo: str, novo: str):
    valor = celula.value
    if isinstance(valor, CellRichText):
        # Monta blocos novos em vez de alterar os existentes: `copy_worksheet`
        # compartilha o mesmo objeto entre as abas copiadas.
        blocos = []
        for bloco in valor:
            if isinstance(bloco, TextBlock):
                texto = bloco.text.replace(antigo, novo)
                if texto:
                    blocos.append(TextBlock(copy.copy(bloco.font), texto))
            elif bloco:
                blocos.append(bloco.replace(antigo, novo))
        celula.value = CellRichText(blocos)
    elif isinstance(valor, str):
        celula.value = valor.replace(antigo, novo)


def _item_protocolo(db: Session, bip: Bipagem) -> dict:
    carga = carga_do_bip(db, bip)
    return {
        "cnpj_filial": carga.cnpj_filial if carga else bip.cnpj_filial,
        # DACTE Nº = conhecimento do parceiro; sem redespacho, o nosso.
        "dacte": _int_se_numero(carga.cte_redespacho) if carga and carga.cte_redespacho else bip.numero,
        "nf": _int_se_numero(carga.notas_fiscais) if carga else None,
        "tomador_cnpj": carga.cnpj_consignatario if carga else None,
        "tomador_nome": (carga.consignatario if carga else None) or "TOMADOR NÃO IDENTIFICADO",
        "tomador_endereco": carga.endereco_consignatario if carga else None,
        "tomador_cidade": carga.cidade_cons if carga else None,
    }


def _copiar_imagens(origem, destino):
    for img in origem._images:
        # `_data()` do openpyxl fecha o buffer da imagem depois de ler --
        # repõe um novo na original, senão o save falha ("closed file").
        dados = img._data()
        img.ref = io.BytesIO(dados)
        nova = Image(io.BytesIO(dados))
        nova.anchor = copy.deepcopy(img.anchor)
        nova.width, nova.height = img.width, img.height
        destino.add_image(nova)


def _titulo_aba(nome: str, usados: set[str]) -> str:
    base = "".join(c for c in nome if c not in "[]:*?/\\")[:28].strip() or "Protocolo"
    titulo, n = base, 2
    while titulo in usados:
        titulo = f"{base[:25]} ({n})"
        n += 1
    usados.add(titulo)
    return titulo


def _recebedor(itens: list[dict]) -> dict:
    """Dados do recebedor de um grupo consolidado (cliente com vários CNPJ,
    mesma raiz): a matriz (/0001) se estiver entre os CT-e, senão o CNPJ
    que mais aparece; o nome é o mais frequente (o mesmo cliente às vezes
    vem como "LTDA" e "LIMITADA")."""
    if not itens:
        return {}
    contagem = Counter(i["tomador_cnpj"] for i in itens if i["tomador_cnpj"])
    matriz = sorted(c for c in contagem if c[8:12] == "0001")
    cnpj = matriz[0] if matriz else (contagem.most_common(1)[0][0] if contagem else None)
    base = next((i for i in itens if i["tomador_cnpj"] == cnpj), itens[0])
    nome = Counter(i["tomador_nome"] for i in itens).most_common(1)[0][0]
    return {**base, "tomador_nome": nome}


def gerar_protocolo_faturas(db: Session, bips: list[Bipagem]) -> bytes:
    # Um grupo (aba) por filial emitente + cliente. Cliente com vários CNPJ
    # (filiais dele, mesma raiz -- ex: FL Brasil, Modular, Capivari) sai
    # consolidado numa aba só.
    grupos: dict[tuple, list[dict]] = {}
    for bip in bips:
        item = _item_protocolo(db, bip)
        if item:
            cliente = item["tomador_cnpj"][:8] if item["tomador_cnpj"] else item["tomador_nome"]
            grupos.setdefault((item["cnpj_filial"], cliente), []).append(item)

    capacidade = len(_SLOTS_PROTOCOLO)
    # (recebedor, itens da aba) -- grupo com mais CT-e do que cabe numa aba
    # continua em outra aba, com o mesmo recebedor.
    folhas = [
        (_recebedor(itens), itens[i:i + capacidade])
        for itens in grupos.values()
        for i in range(0, len(itens), capacidade)
    ] or [({}, [])]

    wb = openpyxl.load_workbook(MODELOS / "protocolo_envio.xlsx", rich_text=True)
    base = wb.active
    # Copia a aba modelo ainda vazia antes de preencher qualquer uma.
    abas = [base]
    for _ in folhas[1:]:
        nova = wb.copy_worksheet(base)
        _copiar_imagens(base, nova)
        nova.page_setup.orientation = base.page_setup.orientation
        nova.page_setup.paperSize = base.page_setup.paperSize
        nova.sheet_properties.pageSetUpPr = copy.copy(base.sheet_properties.pageSetUpPr)
        abas.append(nova)

    titulos: set[str] = set()
    for ws, (primeiro, itens) in zip(abas, folhas):
        filial = dados_filial(primeiro.get("cnpj_filial"))
        razao = filial["razao_social"] or filial["apelido"]

        ws.title = _titulo_aba(primeiro.get("tomador_nome") or "Protocolo", titulos)
        # Recebedor = tomador do serviço.
        ws["B5"] = primeiro.get("tomador_nome")
        ws["B6"] = primeiro.get("tomador_endereco")
        ws["B7"] = fmt_cnpj(primeiro.get("tomador_cnpj")) or None
        ws["E7"] = primeiro.get("tomador_cidade")
        # Emitente (faixa azul do topo): o modelo já vem com os dados da
        # matriz -- troca só o texto dentro da formatação (negrito
        # "CNPJ:"/"Endereço:"), sem perder o rich text da célula.
        _substituir(ws["C1"], _EMITENTE_MODELO["razao_social"], razao)
        _substituir(ws["E1"], _EMITENTE_MODELO["cnpj"], fmt_cnpj(filial["cnpj"]))
        _substituir(ws["C2"], _EMITENTE_MODELO["endereco"], filial["endereco"])

        for (linha, col_dacte, col_nf), item in zip(_SLOTS_PROTOCOLO, itens):
            for coluna, valor in ((col_dacte, item["dacte"]), (col_nf, item["nf"])):
                celula = ws.cell(row=linha, column=coluna, value=valor)
                celula.alignment = celula.alignment.copy(horizontal="center")

    return _salvar(wb)


# ----------------------------------- controle de faturamento (protocolos) ----
#
# Os comprovantes chegam todo dia e são bipados conforme chegam; o protocolo
# de faturas só deve levar os que ainda não foram faturados. Cada protocolo
# gerado marca os bips que entraram nele (`Bipagem.protocolo_id`), e o
# próximo protocolo do mesmo tomador só pega os pendentes.

SEM_TOMADOR = "sem-tomador"


def id_tomador(carga: Carga | None) -> str:
    """Mesmo `id` de `tomadores_disponiveis`: raiz do CNPJ ou "nome:<nome>"."""
    if carga is None:
        return SEM_TOMADOR
    if carga.cnpj_consignatario:
        return carga.cnpj_consignatario[:8]
    if (carga.consignatario or "").strip():
        return PREFIXO_TOMADOR_NOME + carga.consignatario.strip()
    return SEM_TOMADOR


def _pendentes_por_tomador(db: Session, empresas: set[str] | None = None) -> dict[str, dict]:
    grupos: dict[str, dict] = {}
    bips = (
        db.query(Bipagem)
        .filter(Bipagem.protocolo_id.is_(None))
        .order_by(Bipagem.bipado_em.asc())
        .all()
    )
    for bip in bips:
        carga = carga_do_bip(db, bip)
        if not carga_visivel(carga, empresas):
            continue
        tid = id_tomador(carga)
        nome = (carga.consignatario if carga else None) or "TOMADOR NÃO IDENTIFICADO (fora da base)"
        g = grupos.setdefault(tid, {
            "id": tid,
            "nome": nome.strip(),
            "cnpj_raiz": carga.cnpj_consignatario[:8] if carga and carga.cnpj_consignatario else None,
            "quantidade": 0,
            "bips": [],
        })
        g["quantidade"] += 1
        g["bips"].append(bip)
    return grupos


def faturas_pendentes(db: Session, empresas: set[str] | None = None) -> list[dict]:
    """CT-e bipados que ainda não entraram em nenhum protocolo, por tomador
    (com `empresas`, só os CT-e dessas empresas)."""
    grupos = _pendentes_por_tomador(db, empresas)
    return sorted(
        ({k: v for k, v in g.items() if k != "bips"} for g in grupos.values()),
        key=lambda g: (g["nome"] or "").upper(),
    )


def criar_protocolo_faturas(db: Session, tomador: str, empresas: set[str] | None = None) -> ProtocoloFatura | None:
    """Gera o protocolo com os CT-e pendentes do tomador e dá baixa neles.
    Com `empresas` (analista), só leva os CT-e dessas empresas. Retorna
    `None` se não houver pendentes."""
    grupo = _pendentes_por_tomador(db, empresas).get(tomador)
    if not grupo:
        return None

    protocolo = ProtocoloFatura(tomador_id=tomador, tomador_nome=grupo["nome"], quantidade=0)
    db.add(protocolo)
    db.flush()
    # Só marca os que continuam pendentes -- se outra pessoa gerou um
    # protocolo ao mesmo tempo, os que ela já pegou ficam de fora deste.
    ids = [b.id for b in grupo["bips"]]
    db.query(Bipagem).filter(Bipagem.id.in_(ids), Bipagem.protocolo_id.is_(None)).update(
        {Bipagem.protocolo_id: protocolo.id}, synchronize_session=False
    )
    bips = (
        db.query(Bipagem)
        .filter(Bipagem.protocolo_id == protocolo.id)
        .order_by(Bipagem.bipado_em.asc())
        .all()
    )
    if not bips:
        db.rollback()
        return None
    protocolo.quantidade = len(bips)
    protocolo.arquivo = gerar_protocolo_faturas(db, bips)
    db.commit()
    return protocolo


def desfazer_protocolo(db: Session, protocolo_id: int) -> bool:
    """Apaga o protocolo e devolve os CT-e dele pra lista de pendentes."""
    protocolo = db.query(ProtocoloFatura).filter(ProtocoloFatura.id == protocolo_id).first()
    if not protocolo:
        return False
    db.query(Bipagem).filter(Bipagem.protocolo_id == protocolo_id).update(
        {Bipagem.protocolo_id: None}, synchronize_session=False
    )
    db.delete(protocolo)
    db.commit()
    return True
